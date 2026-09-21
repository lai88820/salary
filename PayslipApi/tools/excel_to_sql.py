"""
每月薪資 Excel → MySQL 匯入檔

用法：
    python excel_to_sql.py "11509薪資計算明細.xlsx" --password 你的密碼

會在 Excel 同一個資料夾產生 payroll_115-09.sql，用 MySQL Workbench 開啟後整份執行即可。
同一個月份重複匯入會先刪掉舊的，不會重複。
"""
import argparse, datetime, io, os, re, sys
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

DDL = """SET NAMES utf8mb4;
CREATE DATABASE IF NOT EXISTS `payroll` DEFAULT CHARACTER SET utf8mb4;
USE `payroll`;

CREATE TABLE IF NOT EXISTS `payroll_detail` (
  `id` INT AUTO_INCREMENT PRIMARY KEY,
  `period` VARCHAR(10) NOT NULL COMMENT '薪資月份(民國)',
  `status` VARCHAR(10) NOT NULL COMMENT '在職/離職',
  `department` VARCHAR(50) COMMENT '部門/分校',
  `seq_no` INT COMMENT '編號',
  `emp_code` VARCHAR(20) COMMENT '工號',
  `last_work_day` VARCHAR(20) COMMENT '最後上班日(離職)',
  `name` VARCHAR(50) COMMENT '姓名',
  `english_name` VARCHAR(50) COMMENT '英文名',
  `pay_transfer` TINYINT COMMENT '電匯',
  `pay_cash` TINYINT COMMENT '領現',
  `branch_position` VARCHAR(100) COMMENT '分校/職務',
  `hire_info` VARCHAR(200) COMMENT '到職日',
  `salary` DECIMAL(14,4) COMMENT '薪資',
  `salary_structure` TEXT COMMENT '薪資結構說明',
  `absence_adj` DECIMAL(14,4) COMMENT '未到(加減薪)',
  `absence_note` TEXT,
  `insurance` DECIMAL(14,4) COMMENT '勞健保',
  `insurance_note` TEXT COMMENT '投保說明',
  `other_adj_1` DECIMAL(14,4) COMMENT '其他加減項1',
  `other_adj_1_note` TEXT,
  `other_adj_2` DECIMAL(14,4) COMMENT '其他加減項2',
  `other_adj_2_note` TEXT,
  `total` DECIMAL(14,4) COMMENT '合計/結算',
  `bonus_monthly_quarterly` DECIMAL(14,4) COMMENT '月領/季領感恩獎金',
  `bonus_note` TEXT COMMENT '感恩獎金說明',
  `bonus_3month_note` TEXT COMMENT '3個月領感恩獎金',
  `remarks` TEXT COMMENT '其他備註',
  -- ===== 薪資單（由 Excel 公式拆出，可在資料庫直接修改） =====
  `slip_name` VARCHAR(50) COMMENT '薪資單姓名',
  `slip_dept` VARCHAR(50) COMMENT '薪資單部門(分校)',
  `hire_date` VARCHAR(20) COMMENT '到職日',
  `pay_base` INT COMMENT '本薪',
  `pay_overtime` INT COMMENT '平均加班',
  `pay_other_add` INT COMMENT '其他加項',
  `earn_note` VARCHAR(200) COMMENT '應領備註說明',
  `ded_labor_ins` INT COMMENT '代扣勞保',
  `ded_health_ins` INT COMMENT '代扣健保',
  `ded_pension` INT COMMENT '自提勞退',
  `ded_other` INT COMMENT '其他減項(含未到/遲到扣款)',
  `deduct_note` VARCHAR(200) COMMENT '應扣備註說明',
  `gross_total` INT COMMENT '應領合計(A)',
  `deduct_total` INT COMMENT '應扣合計(B)',
  `net_pay` INT COMMENT '實發金額(A)-(B)',
  KEY `idx_name` (`name`), KEY `idx_emp` (`emp_code`)
) DEFAULT CHARSET=utf8mb4 COMMENT='薪資匯總明細表';

CREATE TABLE IF NOT EXISTS `payroll_class_count` (
  `payroll_id` INT NOT NULL COMMENT '對應 payroll_detail.id',
  `category` VARCHAR(10) NOT NULL COMMENT '安親/資數/輔導資數/兒美/才藝/其他',
  `headcount` DECIMAL(8,1) NOT NULL COMMENT '人數',
  PRIMARY KEY (`payroll_id`, `category`)
) DEFAULT CHARSET=utf8mb4 COMMENT='薪資單 帶班科目人數（由備註自動判讀，請核對）';

CREATE TABLE IF NOT EXISTS `lecturer_payroll` (
  `id` INT AUTO_INCREMENT PRIMARY KEY,
  `period` VARCHAR(10) NOT NULL,
  `subject` VARCHAR(30) COMMENT '科別(英文/數學/外師…)',
  `seq_no` INT NOT NULL COMMENT '編號',
  `emp_code` VARCHAR(20) COMMENT '工號',
  `name` VARCHAR(50),
  `alias` VARCHAR(100) COMMENT '其他名字/英文名',
  `pay_transfer` TINYINT,
  `pay_cash` TINYINT,
  `total` DECIMAL(14,4) COMMENT '合計',
  `remarks` TEXT,
  UNIQUE KEY `uk_seq` (`period`, `seq_no`)
) DEFAULT CHARSET=utf8mb4 COMMENT='講師外師 薪資';

CREATE TABLE IF NOT EXISTS `lecturer_class_line` (
  `id` INT AUTO_INCREMENT PRIMARY KEY,
  `period` VARCHAR(10) NOT NULL,
  `seq_no` INT NOT NULL COMMENT '對應 lecturer_payroll.seq_no',
  `line_no` INT NOT NULL,
  `branch` VARCHAR(50) COMMENT '分校',
  `description` TEXT COMMENT '課程堂數/鐘點說明',
  `absence_adj` DECIMAL(14,4) COMMENT '未到(加減薪)',
  `sessions` DECIMAL(14,4) COMMENT '堂數(1.5H計)',
  `hourly_fee` DECIMAL(14,4) COMMENT '鐘點費(1.5H計)',
  `other_adj` VARCHAR(100) COMMENT '其他加減項',
  KEY `idx_seq` (`period`, `seq_no`)
) DEFAULT CHARSET=utf8mb4 COMMENT='講師每一行課程明細';

CREATE TABLE IF NOT EXISTS `cash_payout` (
  `id` INT AUTO_INCREMENT PRIMARY KEY,
  `period` VARCHAR(10) NOT NULL,
  `payout_type` VARCHAR(20) NOT NULL COMMENT '薪資 / 感恩獎金',
  `payout_date` VARCHAR(10) COMMENT '發放日',
  `category` VARCHAR(20) COMMENT '領現身分',
  `seq_no` INT COMMENT '序號',
  `name` VARCHAR(50),
  `amount` INT NOT NULL,
  `n1000` INT, `n500` INT, `n100` INT, `n50` INT, `n10` INT, `n5` INT, `n1` INT
) DEFAULT CHARSET=utf8mb4 COMMENT='領現金額與鈔票面額';

CREATE TABLE IF NOT EXISTS `salary_split_formula` (
  `role` VARCHAR(20) NOT NULL COMMENT '職務',
  `total_salary` DECIMAL(14,4) NOT NULL COMMENT '總薪資',
  `base_salary` DECIMAL(14,4) COMMENT '本薪',
  `flat_allowance` DECIMAL(14,4) COMMENT '平加',
  PRIMARY KEY (`role`, `total_salary`)
) DEFAULT CHARSET=utf8mb4 COMMENT='本薪&平加公式';

CREATE TABLE IF NOT EXISTS `overtime_by_role` (
  `role` VARCHAR(30) PRIMARY KEY COMMENT '職務',
  `regular_hours` DECIMAL(14,4) COMMENT '不加班時數',
  `ot_equiv_hours` DECIMAL(14,4) COMMENT '加班等價時數',
  `total_equiv_hours` DECIMAL(14,4) COMMENT '總等價時數',
  `base_pct` DECIMAL(10,6) COMMENT '底薪百分比',
  `base_salary` DECIMAL(14,4) COMMENT '底薪',
  `ot_pay` DECIMAL(14,4) COMMENT '加班費',
  `total_pay` DECIMAL(14,4) COMMENT '總共薪資',
  `min_salary` DECIMAL(14,4) COMMENT '最低薪資',
  `hourly_rate` DECIMAL(14,4) COMMENT '時薪',
  `note` VARCHAR(100)
) DEFAULT CHARSET=utf8mb4 COMMENT='各職務加班費計算(新)';"""

# =====================================================================
#  從 Excel 公式拆出 應領 / 應扣 各項
# =====================================================================
import re, math

def rnd(x):  # 四捨五入到元（Excel 方式）
    return int(math.floor(abs(x) + 0.5)) * (1 if x >= 0 else -1)

REF = re.compile(r'\$?([A-Z]{1,2})\$?(\d+)')

def split_terms(expr):
    """把 '=a+b-(c-d)*e' 拆成最外層的 +/- 項目 [(sign, text)]"""
    e = expr.lstrip('=').strip()
    while e.startswith('(') and e.endswith(')') and _balanced(e[1:-1]):
        e = e[1:-1]
    terms, depth, cur, sign = [], 0, '', 1
    for i, ch in enumerate(e):
        if ch == '(':
            depth += 1
        elif ch == ')':
            depth -= 1
        if depth == 0 and ch in '+-' and cur.strip() and cur.strip()[-1] not in '*/':
            terms.append((sign, cur.strip()))
            sign, cur = (1 if ch == '+' else -1), ''
            continue
        if depth == 0 and ch in '+-' and not cur.strip():
            sign = sign * (1 if ch == '+' else -1)
            continue
        cur += ch
    if cur.strip():
        terms.append((sign, cur.strip()))
    return terms

def _balanced(s):
    d = 0
    for ch in s:
        d += ch == '('
        d -= ch == ')'
        if d < 0:
            return False
    return d == 0

def evaluate(text, val):
    py = REF.sub(lambda m: f'({_num(val(m.group(1), int(m.group(2))))})', text.lstrip('='))
    py = py.replace('"', '')
    return float(eval(py, {'__builtins__': {}}))

def _num(v):
    try:
        return float(v) if v not in (None, '', ' ') else 0.0
    except (TypeError, ValueError):
        return 0.0

LABOR_BASES, HEALTH_BASES = set(), set()

def collect_bases(j_formulas):
    """從兩項以上的勞健保公式學習：第 1 項是勞保基數，第 2 項是健保基數"""
    for f in j_formulas:
        if not isinstance(f, str) or not f.startswith('='):
            continue
        t = split_terms(f)
        if len(t) >= 2:
            b1 = re.match(r'(\d+)', t[0][1]); b2 = re.match(r'(\d+)', t[1][1])
            if b1: LABOR_BASES.add(int(b1.group(1)))
            if b2: HEALTH_BASES.add(int(b2.group(1)))

def split_insurance(jf, jv, val):
    """回傳 (勞保, 健保, 其他)，皆為正數的扣款金額"""
    if isinstance(jf, str) and jf.startswith('='):
        terms = split_terms(jf)
        amounts = [-s * evaluate(t, val) for s, t in terms]
        if len(terms) >= 2:
            return amounts[0], amounts[1], sum(amounts[2:])
        base = re.match(r'(\d+)', terms[0][1]) if terms else None
        if base and int(base.group(1)) in HEALTH_BASES and int(base.group(1)) not in LABOR_BASES:
            return 0.0, amounts[0], 0.0
        return amounts[0] if terms else 0.0, 0.0, 0.0
    v = -_num(jv)
    if v and int(round(abs(v))) in HEALTH_BASES and int(round(abs(v))) not in LABOR_BASES:
        return 0.0, v, 0.0
    return v, 0.0, 0.0

STD = re.compile(r'^\(\(?(\d+(?:\.\d+)?)-\$?I(\d+)\)/\1\*\$?H\2\)?$')

def compute(r, f_ws, v_ws, k_note, l_note):
    """r = 主列列號。回傳 dict 或 None（無法解析時）"""
    def val(col, row):
        return v_ws[f'{col}{row}'].value

    total_v = _num(v_ws.cell(r, 13).value)
    mf = f_ws.cell(r, 13).value
    jf, jv = f_ws.cell(r, 10).value, v_ws.cell(r, 10).value

    item = dict(base=0.0, overtime=0.0, other_add=0.0, labor=0.0, health=0.0,
                pension=0.0, other_deduct=0.0)
    ins_done = False

    def classify_adj(amount, note):
        note = note or ''
        if amount >= 0:
            item['overtime' if '平均加班' in note else 'other_add'] += amount
        else:
            item['pension' if '自提' in note else 'other_deduct'] += -amount

    if isinstance(mf, str) and mf.startswith('='):
        for sign, t in split_terms(mf):
            refs = REF.findall(t)
            std = STD.match(t)
            if std and sign == 1:
                n, h, i = float(std.group(1)), _num(val('H', r)), _num(val('I', r))
                item['base'] += h
                item['other_deduct'] += i / n * h  # 未到扣款
                continue
            if len(refs) == 1 and t.replace('$', '') == f'{refs[0][0]}{refs[0][1]}':
                col = refs[0][0]
                amount = sign * _num(val(col, int(refs[0][1])))
                if col == 'J' and not ins_done:
                    lab, hea, oth = split_insurance(jf, jv, val)
                    if sign == -1:  # =H/30-J 這種寫法，J 本身是正數
                        lab, hea, oth = -lab, -hea, -oth
                    item['labor'] += lab; item['health'] += hea; item['other_deduct'] += oth
                    ins_done = True
                    continue
                if col == 'K':
                    classify_adj(amount, k_note); continue
                if col == 'L':
                    classify_adj(amount, l_note); continue
                if col == 'H':
                    item['base'] += amount; continue
            amount = sign * evaluate(t, val)
            if amount >= 0:
                item['base'] += amount  # 時數×時薪、按日計薪 → 本薪
            else:
                item['other_deduct'] += -amount
    else:
        lab, hea, oth = split_insurance(jf, jv, val)
        item.update(labor=lab, health=hea, other_deduct=oth, base=total_v + lab + hea + oth)

    # 四捨五入到元，差額（小數誤差）併入其他減項 / 其他加項，確保 實發 = Excel 合計
    out = {k: rnd(v) for k, v in item.items()}
    gross = out['base'] + out['overtime'] + out['other_add']
    deduct = out['labor'] + out['health'] + out['pension'] + out['other_deduct']
    diff = (gross - deduct) - rnd(total_v)
    if diff > 0:
        out['other_deduct'] += diff
    elif diff < 0:
        out['other_add'] += -diff
    gross = out['base'] + out['overtime'] + out['other_add']
    deduct = out['labor'] + out['health'] + out['pension'] + out['other_deduct']
    out.update(gross=gross, deduct=deduct, net=gross - deduct, raw_diff=diff)
    return out


# ---------------- 帶班人數 ----------------

NUM = r'(\d+(?:\.\d+)?)'

def _current_month_text(text, period='115-08'):
    y, m = period.split('-')
    idx = -1
    for pat in (rf'{y}\s*/\s*0?{int(m)}\s*月?', rf'{int(m)}月'):
        for mt in re.finditer(pat, text):
            idx = max(idx, mt.start())
        if idx >= 0:
            break
    return text[idx:] if idx >= 0 else text

def class_counts(dept, position, notes, period='115-08'):
    text = _current_month_text('\n'.join(n for n in notes if n), period)
    ctx = f'{dept or ""} {position or ""}'
    res = {}

    def last(pattern):
        found = re.findall(pattern, text)
        return float(found[-1]) if found else None

    if '兒美' in ctx:
        v = last(rf'總人數\s*[=＝]\s*{NUM}') or last(rf'共\s*{NUM}\s*人') or last(rf'班\s*[\*＊,，]?\s*{NUM}\s*人')
        if v: res['兒美'] = v
    if '安親' in ctx:
        v = last(rf'安(?:親)?\s*[:：]?\s*(?:小.\s*[A-Z]?\s*)?{NUM}\s*人')
        if v: res['安親'] = v
        v = last(rf'資(?:數)?\s*[:：]?\s*(?:小.\s*[A-Z]?\s*)?{NUM}\s*人')
        if v: res['資數'] = v
    return res


# =====================================================================
#  Excel → SQL 主程式
# =====================================================================

def q(v):
    if v is None: return 'NULL'
    if isinstance(v, bool): return '1' if v else '0'
    if isinstance(v, float): return repr(round(v, 4))
    if isinstance(v, int): return str(v)
    s = str(v).replace('\\', '\\\\').replace("'", "''")
    return "'" + s + "'"

def isnum(v): return isinstance(v, (int, float)) and not isinstance(v, bool)
def clean(v):
    if v is None: return None
    if isinstance(v, datetime.datetime): return f'{v.month}/{v.day}'
    if isinstance(v, str):
        s = v.strip()
        return s or None
    return v
def num(v): return v if isnum(v) else None
def txt(*vals):
    parts = [str(clean(v)) for v in vals if clean(v) is not None and not isnum(clean(v))]
    return '\n'.join(parts) or None
def flat(v):
    v = clean(v)
    return v.replace('\n', ' ') if isinstance(v, str) else v
def flag(v):
    v = clean(v)
    return 1 if isinstance(v, str) and v.lower() == 'v' else 0

def insert_sql(table, cols, rows, batch=100):
    out = []
    for i in range(0, len(rows), batch):
        out.append(f"INSERT INTO `{table}` ({', '.join('`'+c+'`' for c in cols)}) VALUES\n" +
                   ',\n'.join('(' + ', '.join(q(x) for x in r) + ')' for r in rows[i:i+batch]) + ';')
    return out

def last_row(ws, max_cols=20, gap=60):
    """找最後一列有資料的列（避免 max_row 是 1048576 時跑很久）"""
    last, empty = 0, 0
    for r in range(1, min(ws.max_row, 20000) + 1):
        if any(ws.cell(r, c).value not in (None, '', ' ') for c in range(1, max_cols + 1)):
            last, empty = r, 0
        else:
            empty += 1
            if empty > gap and last:
                break
    return last

def payout_date(ws):
    for c in range(1, 15):
        v = clean(ws.cell(1, c).value)
        if isinstance(v, str):
            mt = re.search(r'(\d{1,2}/\d{1,2})\s*發放', v)
            if mt: return mt.group(1)
    return None


def open_workbook(path, password):
    import msoffcrypto
    with open(path, 'rb') as f:
        head = f.read(8)
    if head.startswith(b'PK'):
        data = open(path, 'rb').read()
    else:
        if not password:
            raise SystemExit('這個 Excel 有密碼，請用 --password 指定。')
        office = msoffcrypto.OfficeFile(open(path, 'rb'))
        office.load_key(password=password)
        buf = io.BytesIO()
        try:
            office.decrypt(buf)
        except Exception:
            raise SystemExit('Excel 密碼錯誤，無法開啟。')
        data = buf.getvalue()
    values = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    formulas = openpyxl.load_workbook(io.BytesIO(data))
    return values, formulas


def convert(path, password=None, out_path=None):
    wb, fwb = open_workbook(path, password)
    names = wb.sheetnames
    main = next((n for n in names if re.fullmatch(r'\d{3}\.\d{2}', n.strip())), None)
    if not main:
        raise SystemExit(f'找不到月份工作表（例如 115.09）。目前的工作表：{names}')
    y, m = main.strip().split('.')
    PERIOD = f'{y}-{m}'
    def sheet(suffix):
        return next((n for n in names if n.strip() == f'{main.strip()}{suffix}'), None)

    out = [f'-- 薪資資料 {PERIOD}（來源：{os.path.basename(path)}）',
           f'-- 產生時間：{datetime.datetime.now():%Y-%m-%d %H:%M}',
           '-- 可以重複匯入：會先刪除同月份的舊資料再重新寫入，不影響其他月份。', '',
           DDL, '', 'START TRANSACTION;', '',
           f"DELETE c FROM payroll_class_count c JOIN payroll_detail d ON d.id = c.payroll_id WHERE d.period = '{PERIOD}';",
           f"DELETE FROM payroll_detail WHERE period = '{PERIOD}';",
           f"DELETE FROM lecturer_class_line WHERE period = '{PERIOD}';",
           f"DELETE FROM lecturer_payroll WHERE period = '{PERIOD}';",
           f"DELETE FROM cash_payout WHERE period = '{PERIOD}';", '']
    summary, warnings = [], []

    # ---------- 薪資匯總明細 ----------
    ws, fws = wb[main], fwb[main]
    end = last_row(ws)
    main_rows = [r for r in range(4, end + 1) if isinstance(ws.cell(r, 2).value, int)]
    formula_rows = [r for r in range(4, end + 1) if isinstance(fws.cell(r, 13).value, str) and fws.cell(r, 13).value.startswith('=')]
    if len(main_rows) < len(formula_rows) * 0.5:
        raise SystemExit('Excel 裡的公式沒有計算結果（可能不是用 Excel 存檔的）。\n'
                         '請用 Excel 打開這個檔案，按一下「儲存」後再轉換一次。')
    LABOR_BASES.clear(); HEALTH_BASES.clear()
    collect_bases([fws.cell(r, 10).value for r in main_rows])
    branches = []
    for r in range(4, end + 1):
        a = clean(ws.cell(r, 1).value)
        mt = re.match(r'^\d+\s*(\S+)$', a) if isinstance(a, str) else None
        if mt: branches.append(mt.group(1))
    abs_words = re.compile(r'(事假?|病假?|豪雨|遲到|曠職?|生理|喪假?|公假?|未到|早退)')

    def slip_dept(dept, pos):
        if dept:
            mt = re.match(r'^\d+\s*(\S+)$', dept)
            if mt: return mt.group(1) + '分校'
        for b in branches:
            if b in (pos or '') or b in (dept or ''):
                return b + '分校'
        return (pos or dept or '').split(' ')[0] or None

    def hire_date(h):
        mt = re.search(r'(\d{2,3})\s*[./]\s*(\d{1,2})\s*[./]\s*(\d{1,2})', h or '')
        return f'{int(mt.group(1))}.{int(mt.group(2)):02d}.{int(mt.group(3)):02d}' if mt else None

    def absence_line(notes):
        for n in notes:
            for line in (n or '').split('\n'):
                line = line.strip()
                if abs_words.search(line) and re.search(r'\d+(\.\d+)?\s*(H|h|小時|分|天)', line):
                    return line
        return None

    cols = ['period','status','department','seq_no','emp_code','last_work_day','name','english_name','pay_transfer','pay_cash',
            'branch_position','hire_info','salary','salary_structure','absence_adj','absence_note','insurance','insurance_note',
            'other_adj_1','other_adj_1_note','other_adj_2','other_adj_2_note','total','bonus_monthly_quarterly','bonus_note',
            'bonus_3month_note','remarks',
            'slip_name','slip_dept','hire_date','pay_base','pay_overtime','pay_other_add','earn_note',
            'ded_labor_ins','ded_health_ins','ded_pension','ded_other','deduct_note','gross_total','deduct_total','net_pay']
    out.append(f'-- {main} 薪資匯總明細（每人一筆，帶班人數用 LAST_INSERT_ID 對應）')
    dept = None; status = '在職'; n_staff = 0; n_counts = 0; total_sum = 0.0
    for r in range(4, end + 1):
        a = clean(ws.cell(r, 1).value)
        if isinstance(a, str):
            if '離職' in a: status = '離職'
            if not a.startswith('成大') and '離職' not in a: dept = a.replace('\n', '')
        b = ws.cell(r, 2).value
        if not isinstance(b, int): continue
        mv = [ws.cell(r, c).value for c in range(1, 21)]
        nv = [ws.cell(r + 1, c).value for c in range(1, 21)] if not isinstance(ws.cell(r + 1, 2).value, int) else [None] * 20
        name_parts = str(clean(mv[3]) or '').split('\n')
        eng = clean(nv[3]); eng = eng if isinstance(eng, str) else (name_parts[1].strip() if len(name_parts) > 1 else None)
        cname = name_parts[0].strip()
        remarks = txt(mv[13] if clean(mv[13]) != clean(mv[3]) else None, nv[13] if clean(nv[13]) != eng else None,
                      mv[16], nv[16], mv[17], nv[17], mv[18], nv[18], mv[19], nv[19])
        d = dept if status == '在職' else '離職人員'
        k_note, l_note = txt(mv[10], nv[10]), txt(mv[11], nv[11])

        try:
            calc = compute(r, fws, ws, k_note, l_note)
        except Exception as e:
            warnings.append(f'第 {r} 列 {cname}：公式無法拆解（{e}），本薪先放合計金額，請手動修正')
            t = rnd(num(mv[12]) or 0)
            calc = dict(base=t, overtime=0, other_add=0, labor=0, health=0, pension=0, other_deduct=0, gross=t, deduct=0, net=t)
        if calc['net'] != rnd(num(mv[12]) or 0):
            warnings.append(f'第 {r} 列 {cname}：實發 {calc["net"]} 與 Excel 合計 {rnd(num(mv[12]) or 0)} 不同，請檢查')

        bonus = num(mv[14])
        earn_note = f'另 15 號匯 {rnd(bonus):,}' if bonus else None
        absence = absence_line([k_note, l_note, txt(mv[8], nv[8]), txt(mv[9], nv[9])])
        deduct_note = absence or (f'{y}/{m} 無未到、遲到' if calc['other_deduct'] == 0 else None)

        row = [PERIOD, status, d, b,
               flat(mv[2]) if status == '在職' else None,
               flat(mv[2]) if status == '離職' else None,
               cname, eng, flag(mv[4]), flag(mv[5]), flat(mv[6]),
               flat(nv[6]), num(mv[7]), txt(mv[7], nv[7]),
               num(mv[8]), txt(mv[8], nv[8]) if txt(mv[8], nv[8]) != '未到' else None,
               num(mv[9]), txt(mv[9], nv[9]),
               num(mv[10]), k_note, num(mv[11]), l_note,
               num(mv[12]), bonus, txt(mv[14], nv[14]), txt(mv[15], nv[15]), remarks,
               cname, slip_dept(d if status == '在職' else None, flat(mv[6])), hire_date(txt(nv[6], mv[6])),
               calc['base'], calc['overtime'], calc['other_add'], earn_note,
               calc['labor'], calc['health'], calc['pension'], calc['other_deduct'], deduct_note,
               calc['gross'], calc['deduct'], calc['net']]
        out += insert_sql('payroll_detail', cols, [row])
        counts = class_counts(d, flat(mv[6]), [k_note, l_note, txt(mv[15], nv[15])], PERIOD)
        if counts:
            out.append('SET @pid = LAST_INSERT_ID();')
            out.append('INSERT INTO `payroll_class_count` (`payroll_id`, `category`, `headcount`) VALUES ' +
                       ', '.join(f'(@pid, {q(k)}, {q(float(v))})' for k, v in counts.items()) + ';')
            n_counts += len(counts)
        n_staff += 1
        total_sum += num(mv[12]) or 0
    summary.append(f'薪資匯總明細：{n_staff} 人，合計 {total_sum:,.0f}（帶班人數 {n_counts} 筆）')

    # ---------- 講師外師 ----------
    name = sheet('講師外師')
    if name:
        ws = wb[name]; end = last_row(ws)
        heads, lines = [], []
        subj = None; cur = None; ln = 0
        for r in range(4, end + 1):
            a = clean(ws.cell(r, 1).value)
            if isinstance(a, str): subj = a.replace('\n', '')
            v = [ws.cell(r, c).value for c in range(1, 20)]
            if isinstance(v[1], int):
                nm = clean(v[3])
                cur = [PERIOD, subj, v[1], flat(v[2]), flat(v[3]).split(' ')[0] if flat(v[3]) else None, None,
                       flag(v[4]), flag(v[5]), num(v[12]), None]
                if isinstance(nm, str) and '\n' in nm:
                    cur[4], cur[5] = nm.split('\n', 1)[0], nm.split('\n', 1)[1].replace('\n', ' ')
                heads.append(cur); ln = 0
            elif cur is None:
                continue
            else:
                dd = clean(v[3])
                if isinstance(dd, str) and dd != cur[4]:
                    cur[5] = (cur[5] + ' / ' if cur[5] else '') + dd
                extra = txt(*v[13:19])
                if extra and extra != cur[4]:
                    cur[9] = (cur[9] + '\n' if cur[9] else '') + extra
            if any(clean(x) is not None for x in (v[6], v[7], v[8], v[9], v[10], v[11])):
                ln += 1
                oth = clean(v[11])
                lines.append([PERIOD, cur[2], ln, flat(v[6]), clean(v[7]) if not isnum(clean(v[7])) else str(v[7]),
                              num(v[8]), num(v[9]), num(v[10]), None if oth is None else str(oth)])
        out += ['', f'-- {name}']
        out += insert_sql('lecturer_payroll', ['period','subject','seq_no','emp_code','name','alias','pay_transfer','pay_cash','total','remarks'], heads)
        out += insert_sql('lecturer_class_line', ['period','seq_no','line_no','branch','description','absence_adj','sessions','hourly_fee','other_adj'], lines)
        summary.append(f'講師外師：{len(heads)} 人，合計 {sum(h[8] or 0 for h in heads):,.0f}')
    else:
        warnings.append(f'找不到「{main}講師外師」工作表，略過')

    # ---------- 算現金 ----------
    cash = []
    for suffix, ptype in (('算現金', '薪資'), ('算現金 感恩獎金', '感恩獎金')):
        name = sheet(suffix)
        if not name:
            continue
        ws = wb[name]; cat = None; pdate = payout_date(ws)
        for r in range(2, last_row(ws) + 1):
            k = clean(ws.cell(r, 11).value)
            if k in ('個數', '面額', '總金額'): break
            if k: cat = str(k).replace('\n', '')
            amt = ws.cell(r, 3).value; nm = clean(ws.cell(r, 12).value)
            if not isnum(amt) or nm is None: continue
            seq = ws.cell(r, 2).value if ptype == '感恩獎金' else None
            cash.append([PERIOD, ptype, pdate, cat, seq if isinstance(seq, int) else None, nm, amt] +
                        [ws.cell(r, c).value for c in range(4, 11)])
    if cash:
        out += ['', '-- 領現金額']
        out += insert_sql('cash_payout', ['period','payout_type','payout_date','category','seq_no','name','amount','n1000','n500','n100','n50','n10','n5','n1'], cash)
        summary.append(f'領現：{len(cash)} 筆，合計 {sum(c[6] for c in cash):,.0f}')

    # ---------- 參考表（非每月資料，有的話整份更新） ----------
    if '本薪&平加公式' in names:
        ws = wb['本薪&平加公式']; sp = []
        for base in (1, 5, 9):
            role = clean(ws.cell(2, base).value)
            for r in range(2, 40):
                t = ws.cell(r, base + 1).value
                if isnum(t): sp.append([role, t, num(ws.cell(r, base + 2).value), num(ws.cell(r, base + 3).value)])
        out += ['', '-- 本薪&平加公式', 'DELETE FROM salary_split_formula;']
        out += insert_sql('salary_split_formula', ['role','total_salary','base_salary','flat_allowance'], sp)
    if '各職務加班費計算 (新)' in names:
        ws = wb['各職務加班費計算 (新)']; ot = []
        for r in range(1, 30):
            v = [ws.cell(r, c).value for c in range(1, 12)]
            if clean(v[0]) and isnum(v[1]):
                ot.append([flat(v[0])] + [num(x) for x in v[1:10]] + [txt(v[10]) if not isnum(v[10]) else str(v[10])])
        out += ['', '-- 各職務加班費計算', 'DELETE FROM overtime_by_role;']
        out += insert_sql('overtime_by_role', ['role','regular_hours','ot_equiv_hours','total_equiv_hours','base_pct','base_salary','ot_pay','total_pay','min_salary','hourly_rate','note'], ot)

    out += ['', 'COMMIT;', '']
    out_path = out_path or os.path.join(os.path.dirname(os.path.abspath(path)), f'payroll_{PERIOD}.sql')
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(out))
    return PERIOD, out_path, summary, warnings


if __name__ == '__main__':
    ap = argparse.ArgumentParser(description='把每月薪資計算明細 Excel 轉成 MySQL 匯入檔')
    ap.add_argument('excel', help='Excel 檔案路徑')
    ap.add_argument('--password', '-p', help='Excel 開啟密碼')
    ap.add_argument('--out', '-o', help='輸出的 .sql 路徑（預設與 Excel 同資料夾）')
    args = ap.parse_args()
    period, path, summary, warnings = convert(args.excel, args.password, args.out)
    print(f'\n月份：{period}')
    for s in summary: print('  ' + s)
    if warnings:
        print(f'\n注意（{len(warnings)} 項）：')
        for w in warnings: print('  - ' + w)
    print(f'\n已產生：{path}')
